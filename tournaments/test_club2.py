"""Tests for the second batch of club features (October 2026)."""
import datetime
import io
import zipfile
from datetime import timedelta

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone, translation

from players.models import DivisionCategory, Player, Team

from .court_access import court_path, ensure_key
from .models import Division, Match, MedalOverride, Registration, Tournament, TournamentDay
from .scheduler import generate_schedule
from .tests import make_user

LOCMEM = {'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 'club2-tests'}}


def _player(owner, name, gender='M', group='', school='', arrived=True):
    return Player.objects.create(owner=owner, name=name, gender=gender, division=group, school=school,
                                 has_arrived=arrived)


def _single(p):
    return Team.objects.create(player1=p)


def _complete(m, winner_is_t1=True, score=None):
    m.score = score or ('15-10, 15-8' if winner_is_t1 else '10-15, 8-15')
    m.winner = m.team1 if winner_is_t1 else m.team2
    m.status = 'completed'
    m.save()


class BaseTournament(TestCase):
    def setUp(self):
        self.user = make_user(username='clubtest2')
        self.c = Client()
        self.c.force_login(self.user)
        self.t = Tournament.objects.create(name="Tardor", date=timezone.localdate(), division_model='youth',
                                           scoring_model='best_of_5_15', owner=self.user)
        TournamentDay.objects.create(tournament=self.t, day_number=1, date=self.t.date,
                                     start_time=datetime.time(0, 5), court_count=3, buffer_minutes=0)


class SmallFixesTest(BaseTournament):
    def test_delete_accepts_catalan_and_legacy_phrases(self):
        url = reverse('tournament_delete', args=[self.t.pk])
        self.c.post(url, {'confirm': 'wrong'})
        self.assertTrue(Tournament.objects.filter(pk=self.t.pk).exists())
        self.c.post(url, {'confirm': '  slet   turnering '})
        self.assertFalse(Tournament.objects.filter(pk=self.t.pk).exists())

    def test_delete_page_shows_translated_phrase(self):
        with translation.override('ca'):
            from django.utils.translation import gettext as _
            phrase = _("DELETE TOURNAMENT")
        self.c.cookies['django_language'] = 'ca'
        r = self.c.get(reverse('tournament_delete', args=[self.t.pk]))
        self.assertContains(r, phrase)
        self.c.post(reverse('tournament_delete', args=[self.t.pk]), {'confirm': phrase.lower()})
        self.assertFalse(Tournament.objects.filter(pk=self.t.pk).exists())

    def test_admin_user_form_has_club_name_and_owner_labels(self):
        admin = User.objects.create_superuser('boss', 'b@example.org', 'x-pass-123456')
        self.user.first_name = 'Club Bàdminton Igualada'
        self.user.save()
        c = Client()
        c.force_login(admin)
        r = c.get('/admin/auth/user/add/')
        self.assertContains(r, 'name="first_name"')
        r = c.post('/admin/auth/user/add/', {'username': 'nouclub', 'first_name': 'Club Nou',
                                             'password1': 'Uns4fe-pass-99', 'password2': 'Uns4fe-pass-99',
                                             'usable_password': 'true'})
        self.assertEqual(User.objects.get(username='nouclub').first_name, 'Club Nou')
        r = c.get('/admin/players/divisioncategory/add/')
        self.assertContains(r, 'Club Bàdminton Igualada (clubtest2)')


class SchoolTest(BaseTournament):
    def test_form_reuses_existing_spelling(self):
        _player(self.user, 'A', school='Escola Pia')
        r = self.c.post(reverse('player_add'), {'name': 'B', 'gender': 'M', 'division': '', 'school': ' escola  pia '})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Player.objects.get(name='B').school, 'Escola Pia')

    def test_excel_upload_with_catalan_headers(self):
        DivisionCategory.objects.create(owner=self.user, name='Aleví')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Nom', 'Sexe', 'Edat', "Grup d'edat", 'Escola / club'])
        ws.append(['Anna Puig', 'F', 11, 'aleví', 'Escola Pia'])
        ws.append(['Pol Roca', 'M', '', '', ''])
        buf = io.BytesIO()
        wb.save(buf)
        f = SimpleUploadedFile('j.xlsx', buf.getvalue())
        self.c.post(reverse('player_upload'), {'excel_file': f})
        anna = Player.objects.get(name='Anna Puig')
        self.assertEqual((anna.gender, anna.division, anna.school), ('K', 'Aleví', 'Escola Pia'))
        self.assertEqual(Player.objects.get(name='Pol Roca').school, '')

    def test_template_download_is_a_valid_workbook(self):
        DivisionCategory.objects.create(owner=self.user, name='Benjamí')
        r = self.c.get(reverse('player_template_download'))
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.active.max_column, 5)

    def test_school_shown_next_to_names_and_links(self):
        a, b = _player(self.user, 'Anna', school='Escola Pia'), _player(self.user, 'Bru')
        d = Division.objects.create(tournament=self.t, name='Aleví', discipline='single', tournament_type='group')
        d.teams.set([_single(a), _single(b)])
        generate_schedule(d)
        r = Client().get(reverse('public_tournament', args=[self.t.pk]))
        self.assertContains(r, 'Escola Pia')
        self.assertContains(r, reverse('public_player', args=[self.t.pk, a.pk]))

    def test_export_import_keeps_school(self):
        a = _player(self.user, 'Anna', school='Escola Pia')
        d = Division.objects.create(tournament=self.t, name='Aleví', discipline='single', tournament_type='group',
                                    medals_by_age_group=True)
        d.teams.set([_single(a), _single(_player(self.user, 'Bru'))])
        from .views import tournament_export_data
        data = tournament_export_data(self.t)
        self.assertEqual(data['players'][0]['school'] if data['players'][0]['name'] == 'Anna'
                         else data['players'][1]['school'], 'Escola Pia')
        self.assertTrue(data['divisions'][0]['medals_by_age_group'])


class PublicPlayerTest(BaseTournament):
    def test_player_page(self):
        a, b = _player(self.user, 'Anna'), _player(self.user, 'Bru')
        d = Division.objects.create(tournament=self.t, name='Aleví', discipline='single', tournament_type='group')
        d.teams.set([_single(a), _single(b)])
        generate_schedule(d)
        _complete(d.matches.first(), winner_is_t1=d.matches.first().team1.player1 == a)
        r = Client().get(reverse('public_player', args=[self.t.pk, a.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Anna')
        self.assertContains(r, 'Bru')
        self.assertEqual(r.context['won_total'], 1)
        # first set is shown from Anna's side
        self.assertTrue(r.context['played'][0].score_own.startswith('15-'))
        outsider = _player(self.user, 'Carla')
        self.assertEqual(Client().get(reverse('public_player', args=[self.t.pk, outsider.pk])).status_code, 404)


class MedalsByAgeGroupTest(BaseTournament):
    def test_three_medals_per_age_group_and_gender(self):
        with translation.override('en'):
            self._check_medals()

    def _check_medals(self):
        names = [('Anna', 'K', 'Benjamí'), ('Berta', 'K', 'Benjamí'), ('Clara', 'K', 'Benjamí'), ('Dana', 'K', 'Benjamí'),
                 ('Eloi', 'M', 'Benjamí'), ('Ferran', 'M', 'Benjamí'), ('Gina', 'K', 'Aleví'), ('Hug', 'M', 'Aleví')]
        teams = [_single(_player(self.user, n, g, grp)) for n, g, grp in names]
        d = Division.objects.create(tournament=self.t, name='Mixta', discipline='single', tournament_type='groups',
                                    group_count=2, medals_by_age_group=True)
        d.teams.set(teams)
        generate_schedule(d)
        order = {t.pk: i for i, t in enumerate(teams)}
        for m in d.matches.all():
            _complete(m, winner_is_t1=order[m.team1_id] < order[m.team2_id])
        from .medals import compute_subgroup_medals
        sgs = {sg['label']: sg for sg in compute_subgroup_medals(d)}
        self.assertIn('Benjamí · girls', sgs)
        self.assertIn('Aleví · boys', sgs)
        girls = sgs['Benjamí · girls']
        self.assertEqual(len(girls['gold'] + girls['silver'] + girls['bronze']), 3)
        self.assertEqual(girls['gold'][0].name, 'Anna')
        # manual correction for one subgroup only
        MedalOverride.objects.create(division=d, subgroup='Benjamí|K', medal='gold', team=teams[3], order=1)
        sgs = {sg['label']: sg for sg in compute_subgroup_medals(d)}
        self.assertEqual(sgs['Benjamí · girls']['gold'][0].name, 'Dana')
        self.assertEqual(sgs['Benjamí · boys']['gold'][0].name, 'Eloi')
        r = self.c.get(reverse('division_medals', args=[d.pk]))
        self.assertContains(r, 'Benjamí · girls')
        r = self.c.get(reverse('division_medals_edit', args=[d.pk]))
        self.assertEqual(r.status_code, 200)

    def test_mode_toggle(self):
        d = Division.objects.create(tournament=self.t, name='X', discipline='single', tournament_type='group')
        self.c.post(reverse('division_medals_mode', args=[d.pk]), {'by_age_group': '1'})
        d.refresh_from_db()
        self.assertTrue(d.medals_by_age_group)


class LastMinuteTest(BaseTournament):
    def setUp(self):
        super().setUp()
        self.ps = [_player(self.user, n) for n in ('Anna', 'Bru', 'Carla', 'Dani')]
        self.d = Division.objects.create(tournament=self.t, name='Aleví', discipline='single', tournament_type='group')
        self.d.teams.set([_single(p) for p in self.ps])
        generate_schedule(self.d)

    def test_walkover_score_is_written_from_team1_side(self):
        from .live_ops import record_walkover
        m = self.d.matches.first()
        record_walkover(m, m.team2)
        m.refresh_from_db()
        self.assertEqual((m.score, m.winner_id, m.status, m.walkover), ('0-15, 0-15', m.team2_id, 'completed', True))
        self.assertIsNotNone(m.finished_at)

    def test_withdraw_gives_walkovers(self):
        r = self.c.post(reverse('tournament_player_withdraw', args=[self.t.pk, self.ps[0].pk]))
        self.assertEqual(r.status_code, 302)
        anna_matches = Match.objects.filter(division=self.d, team1__player1=self.ps[0]) | \
            Match.objects.filter(division=self.d, team2__player1=self.ps[0])
        self.assertTrue(all(m.status == 'completed' and m.walkover for m in anna_matches))
        self.assertTrue(all(m.winner.player1 != self.ps[0] for m in anna_matches))
        self.assertTrue(Player.objects.get(pk=self.ps[0].pk).withdrawn)
        self.c.post(reverse('tournament_player_reinstate', args=[self.t.pk, self.ps[0].pk]))
        self.assertFalse(Player.objects.get(pk=self.ps[0].pk).withdrawn)

    def test_court_referee_walkover(self):
        self.t.court_access_enabled = True
        self.t.save()
        ensure_key(self.t)
        m = Match.objects.filter(division=self.d).first()
        url = court_path(self.t, 1).replace('/pista/', '/pista/')
        sig = url.rstrip('/').split('/')[-1]
        r = Client().post(reverse('court_walkover', args=[self.t.pk, 1, sig, m.pk]), {'winner': m.team1_id})
        self.assertEqual(r.status_code, 302)
        m.refresh_from_db()
        self.assertEqual((m.status, m.walkover, m.winner_id, m.entered_by_court), ('completed', True, m.team1_id, '1'))

    def test_replan_from_now(self):
        from .schedule_planner import generate_time_schedule, replan_pending
        for i, m in enumerate(Match.objects.filter(division=self.d).order_by('pk'), start=1):
            m.match_number = i
            m.save(update_fields=['match_number'])
        self.assertEqual(generate_time_schedule(self.t), 6)
        first = Match.objects.filter(division=self.d).order_by('scheduled_time').first()
        first.status, first.court, first.started_at = 'in_progress', '1', timezone.now()
        first.save()
        n = replan_pending(self.t)
        self.assertEqual(n, Match.objects.filter(division=self.d, status='pending').count())
        now = timezone.now()
        for m in Match.objects.filter(division=self.d, status='pending'):
            self.assertGreaterEqual(m.scheduled_time, now - timedelta(minutes=1))
        r = self.c.post(reverse('tournament_replan', args=[self.t.pk]))
        self.assertEqual(r.status_code, 302)


class ReceptionCallsDelayTest(BaseTournament):
    def setUp(self):
        super().setUp()
        self.ps = [_player(self.user, n, arrived=False) for n in ('Anna', 'Bru', 'Carla', 'Dani')]
        self.d = Division.objects.create(tournament=self.t, name='Aleví', discipline='single', tournament_type='group')
        self.d.teams.set([_single(p) for p in self.ps])
        generate_schedule(self.d)

    def test_reception_page(self):
        r = self.c.get(reverse('tournament_reception', args=[self.t.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual((r.context['present'], r.context['absent']), (0, 4))

    def test_calls_for_free_courts_and_mark_called(self):
        from .live_info import compute_calls
        calls = compute_calls(self.t)
        self.assertEqual(len([c for c in calls if c['reason'] == 'court_free']), 2)  # 4 players → only 2 matches at once
        m = calls[0]['match']
        self.c.post(reverse('match_mark_called', args=[m.pk]))
        m.refresh_from_db()
        self.assertIsNotNone(m.called_at)
        m.status, m.court, m.started_at = 'in_progress', '1', timezone.now()
        m.save()
        self.assertNotIn(m.pk, [c['match'].pk for c in compute_calls(self.t)])
        r = self.c.get(reverse('tournament_calls', args=[self.t.pk]))
        self.assertEqual(r.status_code, 200)

    def test_delay(self):
        from .live_info import compute_delay
        now = timezone.now()
        self.assertEqual(compute_delay(self.t, now), 0)
        for i, m in enumerate(self.d.matches.all()[:3]):
            m.scheduled_time = now - timedelta(minutes=60 - i * 10)
            m.started_at = m.scheduled_time + timedelta(minutes=20)
            m.status = 'completed'
            m.save()
        self.assertEqual(compute_delay(self.t, now), 20)


@override_settings(CACHES=LOCMEM)
class FinalAndRegistrationTest(BaseTournament):
    def setUp(self):
        super().setUp()
        a, b = _player(self.user, 'Anna', 'K', school='Escola Pia'), _player(self.user, 'Bru')
        self.d = Division.objects.create(tournament=self.t, name='Aleví', discipline='single', tournament_type='group',
                                         gold_count=1, silver_count=1)
        self.d.teams.set([_single(a), _single(b)])
        generate_schedule(self.d)
        for m in self.d.matches.all():
            m.started_at = timezone.now() - timedelta(minutes=20)
            _complete(m)

    def test_medal_table_and_diplomas(self):
        r = self.c.get(reverse('tournament_medal_table', args=[self.t.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['summary']['participants'], 2)
        self.assertEqual(r.context['summary']['schools'], 1)
        r = self.c.get(reverse('tournament_diplomas', args=[self.t.pk]))
        self.assertEqual(len(r.context['diplomas']), 2)
        r = self.c.get(reverse('tournament_diplomas', args=[self.t.pk]) + '?all=1')
        self.assertEqual(len(r.context['diplomas']), 2)

    def test_public_registration_flow(self):
        url = reverse('public_register', args=[self.t.pk])
        r = Client().get(url)
        self.assertContains(r, '🔒')
        self.t.registration_open = True
        self.t.save()
        anon = Client()
        r = anon.post(url, {'name': 'Nil  Soler', 'gender': 'M', 'school': 'escola pia', 'contact': '600000000',
                            'division': self.d.pk, 'consent': 'on'})
        self.assertEqual(r.status_code, 302)
        reg = Registration.objects.get()
        self.assertEqual((reg.name, reg.school, reg.status), ('Nil Soler', 'Escola Pia', 'pending'))
        # honeypot: silently ignored
        anon.post(url, {'name': 'Bot', 'gender': 'M', 'consent': 'on', 'website': 'spam'})
        self.assertEqual(Registration.objects.count(), 1)
        r = self.c.get(reverse('tournament_registrations', args=[self.t.pk]))
        self.assertContains(r, 'Nil Soler')
        self.c.post(reverse('registration_accept', args=[self.t.pk, reg.pk]), {'division': self.d.pk})
        reg.refresh_from_db()
        self.assertEqual(reg.status, 'accepted')
        self.assertTrue(self.d.teams.filter(player1__name='Nil Soler').exists())

    def test_backup_zip_and_connection_test(self):
        r = self.c.get(reverse('club_backup'))
        z = zipfile.ZipFile(io.BytesIO(r.content))
        self.assertIn('jugadors.json', z.namelist())
        self.assertTrue(any(n.startswith('tornejos/') for n in z.namelist()))
        self.assertEqual(Client().get(reverse('connection_test')).status_code, 200)
        self.assertEqual(len(Client().get(reverse('connection_payload')).content), 256 * 1024)
