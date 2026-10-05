"""Excel template for registering players (filled in by schools/clubs, then uploaded)."""
from io import BytesIO

from django.utils.translation import gettext as _

from .models import DivisionCategory


def template_headers():
    """Column titles in the active language, in template order, with their field name."""
    return [
        ('name', _("Name")),
        ('gender', _("Gender")),
        ('age', _("Age")),
        ('division', _("Age group")),
        ('school', _("School / club")),
    ]


def build_player_template(owner):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _("Players")[:31]
    head_fill = PatternFill('solid', fgColor='111111')
    head_font = Font(bold=True, color='FFFFFF')
    widths = [32, 10, 8, 18, 30]
    for i, ((_f, label), w) in enumerate(zip(template_headers(), widths), start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    # Drop-down lists: gender and the club's age groups
    dv = DataValidation(type='list', formula1='"M,F"', allow_blank=False)
    dv.error = _("Write M (male) or F (female).")
    ws.add_data_validation(dv)
    dv.add('B2:B500')
    groups = list(DivisionCategory.objects.filter(owner=owner).values_list('name', flat=True))
    if groups:
        lists = wb.create_sheet(_("Lists")[:31])
        for r, g in enumerate(groups, start=1):
            lists.cell(row=r, column=1, value=g)
        lists.sheet_state = 'hidden'
        dv2 = DataValidation(type='list', formula1=f"='{lists.title}'!$A$1:$A${len(groups)}", allow_blank=True)
        ws.add_data_validation(dv2)
        dv2.add('D2:D500')

    help_ws = wb.create_sheet(_("Instructions")[:31])
    lines = [
        _("How to fill in this template"),
        "",
        _("One row per player. Do not change the titles of the first row."),
        _("Name: first name and surnames (required)."),
        _("Gender: M (male) or F (female) (required)."),
        _("Age: optional."),
        _("Age group: choose it from the list (e.g. Benjamí, Aleví…)."),
        _("School / club: the school or club the player represents. Leave it empty if they play independently."),
        "",
        _("Then upload the file from Players → Upload Excel."),
    ]
    for r, line in enumerate(lines, start=1):
        c = help_ws.cell(row=r, column=1, value=line)
        if r == 1:
            c.font = Font(bold=True, size=13, color='EB1C60')
    help_ws.column_dimensions['A'].width = 100

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
